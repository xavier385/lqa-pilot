"""Regression fixtures: generic Excel, large streaming upload, launch ordering. No live GPT/ADB."""
import http.client
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
from urllib.parse import quote
from zipfile import ZipFile,ZIP_STORED
from openpyxl import Workbook,load_workbook
from lqa_pilot.android import Android,validate_action
from lqa_pilot.web_bridge import Companion
from lqa_pilot.workbook_project import inspect_workbook,import_project
from lqa_pilot.workbook_understanding import fallback_template
from lqa_pilot.util import read_json,PilotError
from test_adaptive_workbooks import Planner,spec


class GenericProjectsTests(unittest.TestCase):
    def test_package_metadata_uses_default_namespaces(self):
        import xml.etree.ElementTree as ET
        from lqa_pilot.client_excel import xml,tag,C,P
        for namespace,name in [(C,'Types'),(P,'Relationships')]:
            serialized=xml(ET.Element(tag(namespace,name)))
            self.assertIn(('<'+name+' xmlns=').encode(),serialized)

    def test_terms_require_explicit_operator_option_and_exclude_optional_ads(self):
        from lqa_pilot.demo import action
        step=action('tap','Nutzungsvereinbarung checkbox','required_terms',.1,.2)
        with self.assertRaises(PilotError):validate_action(step)
        validate_action(step,allow_required_terms=True)
        for target in ['Allen zustimmen','Personalized advertising checkbox']:
            with self.assertRaises(PilotError):validate_action({**step,'target':target},allow_required_terms=True)

    def test_objective_is_the_only_required_workbook_data(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);w=Workbook();w.active.append(['GOAL','Inspect tutorial and menus']);w.save(root/'input.xlsx')
            p=read_json(import_project(root/'input.xlsx',root/'project',{'language':'German','game_name':'Example Game'},model=Planner()))
            self.assertTrue(p['import_summary']['ready']);self.assertEqual(p['device']['game_name'],'Example Game')
            self.assertFalse(p['glossary']);self.assertFalse(p['project_context']['dev_keys'])
            out=load_workbook(root/'project/client-template.xlsx')
            self.assertEqual([c.value for c in out.active[1]],['Objective / Location','Result','Screenshot','Error type','Text','Description','Suggested fix'])
            out.close()

    def test_unusable_optional_report_falls_back_instead_of_blocking(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);w=Workbook();w.active.append(['GOAL','Inspect menus']);w.active.append(['DEFECTS']);w.save(root/'input.xlsx')
            p=read_json(import_project(root/'input.xlsx',root/'p',{'language':'German','game_name':'Example'},model=Planner([spec('Missing sheet')])))
            self.assertTrue(p['import_summary']['ready']);self.assertTrue(p['import_summary']['fallback_report'])
            self.assertTrue(p['import_summary']['report_warnings'])

    def test_large_file_and_sparse_dimensions_have_no_old_caps(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);file=root/'large.xlsx';w=Workbook();w.active.cell(50000,250,'Objective at a distant cell')
            for i in range(41):w.create_sheet('Sheet'+str(i))
            w.save(file)
            with ZipFile(file,'a') as z:
                with z.open('unused-large-part.bin','w') as part:
                    for _ in range(31):part.write(b'0'*(1024*1024))
            self.assertGreater(file.stat().st_size,30*1024*1024)
            snapshot=inspect_workbook(file,root)
            self.assertEqual(len(snapshot['sheets']),42)
            self.assertEqual(snapshot['sheets'][0]['rows'][0]['cells']['IP'],'Objective at a distant cell')

    def test_http_large_upload_is_spooled_to_disk_before_worker(self):
        with tempfile.TemporaryDirectory() as t:
            c=Companion(t,'https://qa.example.com',port=0);c.start();token=c.pair(c.origin,c.code)
            received=[]
            def submit(path,name,options):
                self.assertIsInstance(path,Path);received.append(path.stat().st_size);return {'status':'importing'}
            client=http.client.HTTPConnection('127.0.0.1',c.port,timeout=15)
            try:
                size=31*1024*1024
                with patch.object(c.jobs,'submit',side_effect=submit):
                    client.putrequest('POST','/jobs');client.putheader('Origin',c.origin);client.putheader('Authorization','Bearer '+token)
                    client.putheader('Content-Length',str(size));client.putheader('X-Filename','large.xlsx')
                    client.putheader('X-Options',quote(json.dumps({'mode':'fast','language':'German','game_name':'Example'})));client.endheaders()
                    for _ in range(31):client.send(b'0'*(1024*1024))
                    r=client.getresponse();self.assertEqual(r.status,202);r.read()
                self.assertEqual(received,[size]);self.assertFalse(list(Path(t).glob('upload-*')))
            finally:client.close();c.server.shutdown();c.server.server_close()

    def driver(self,config):
        with patch('lqa_pilot.android.executable',return_value='adb'):return Android(config)

    def test_package_prepare_stops_only_target_without_clearing_data(self):
        d=self.driver({'package':'com.example.game'});d.adb=Mock(return_value=b'package:/data/app/base.apk')
        d.prepare_launch();d.adb.assert_any_call('shell','am','force-stop','com.example.game')
        self.assertFalse(any('clear' in c.args for c in d.adb.call_args_list))

    def test_name_preparation_does_not_tap_until_launch(self):
        d=self.driver({'game_name':'Last Asylum: Plague'});d.foreground=Mock(return_value='com.mumu.launcher')
        raw=b'<hierarchy><node text="Last Asylum" bounds="[10,20][100,110]"/></hierarchy>'
        def adb(*args,**kwargs):
            if args[:4]==('shell','pm','list','packages'):return b'package:com.demo.game\n'
            if args[:2]==('shell','cat'):return raw
            return b''
        d.adb=Mock(side_effect=adb)
        with patch('lqa_pilot.android.time.sleep'):d.prepare_launch()
        self.assertFalse(any('tap' in c.args for c in d.adb.call_args_list))
        d.foreground=Mock(side_effect=['com.mumu.launcher','com.demo.game']);d.launch()
        self.assertEqual(d.package,'com.demo.game');d.adb.assert_any_call('shell','input','tap','55','65')

    def test_recording_is_active_before_launch_and_gpt(self):
        from lqa_pilot.engine import Engine
        from lqa_pilot.demo import DemoAndroid
        from test_modes import project,VideoModel,FakeRecorder
        events=[]
        class Driver(DemoAndroid):
            def prepare_launch(self):events.append('prepare')
            def launch(self):events.append('launch')
        class Recorder(FakeRecorder):
            def start(self):events.append('record');super().start()
            def perform(self,operation):return operation()
        class Model(VideoModel):
            def ask(self,*args,**kwargs):events.append('model');return super().ask(*args,**kwargs)
        with tempfile.TemporaryDirectory() as t:
            root=Path(t)
            with patch('lqa_pilot.video.CaseRecorder',Recorder),patch('lqa_pilot.engine.Path.home',return_value=root),patch('lqa_pilot.engine.time.sleep'):
                state=Engine(project(root,'video'),root/'run',driver=Driver(),model=Model()).run()
            self.assertEqual(events[:4],['prepare','record','launch','model'])
            self.assertEqual(state['cases']['DEMO-BUG']['status'],'RECORDED')
