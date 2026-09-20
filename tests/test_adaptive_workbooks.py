"""Independent spreadsheet layouts, fake planning only. No Android or model calls."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from openpyxl.comments import Comment
from PIL import Image

from lqa_pilot.workbook_project import import_project,inspect_workbook
from lqa_pilot.workbook_understanding import validate_reports
from lqa_pilot.client_excel import export_client
from lqa_pilot.util import PilotError,read_json
from lqa_pilot.web_bridge import Jobs


def spec(name='自由形式',start=71,span=1):
    return {'id':'text','name':name,'purpose':'Textual defects','header_row':start-1,'data_start':start,
        'data_end':start+span-1,'style_row':start,'row_span':span,
        'columns':[{'column':2,'row_offset':0,'field':'description','literal':''}]}


class Planner:
    def __init__(self,reports=None,blocked=False,missing=False):
        self.reports=reports or [];self.blocked=blocked;self.missing=missing;self.layout_input=None;self.context_input=None;self.seen=[]
    def ask(self,role,prompt,schema,images=()):
        self.before_call(role);data=json.loads(prompt.split('DATA:\n')[1])
        if prompt.startswith('Map output'):
            self.layout_input=data
            return {'reports':copy.deepcopy(self.reports),'notes':[],'blockers':[]}
        if prompt.startswith('Review the plan'):
            self.context_input=data
            return {'ready':not self.blocked,'issues':['Essential guide unavailable'] if self.blocked else [],'context_summary':'Client-specific rules',
                'cases':[{'id':c['id'],'guide':'Open messages using the supplied shortcut, then read all tabs.',
                    'dev_keys':[k['id'] for k in data['dev_keys']],'log_keys':[k['id'] for k in data['log_keys']]} for c in data['cases']]}
        result={'rows':[],'cases':[],'instructions':[],'glossary':[],'cheats':[],'dev_keys':[],'log_keys':[],
                'report_headers':[],'taxonomy':[],'notes':[],'continuation':'Mixed project sections'}
        for row in data['rows']:
            n=row['row'];self.seen.append((data['sheet'],n));text=row['cells'].get('A','');ids=[];roles=['context']
            if text=='GOAL':
                ids=['client-goal'];roles=['cases'];result['cases'].append({'id':ids[0],'title':'Messages','objective':'Read every tab',
                    'journey':['Find the messages menu'],'checks':[{'id':'read','scope':'All tabs','instruction':'Check visible text','modality':'visual'}],'source_rows':[n]})
            if text=='GUIDE':result['instructions'].append(row['cells']['B']);roles=['guides']
            if text=='TERM':
                roles=['glossary'];result['glossary'].append({'source':'Inbox','target':'Posteingang','scope':'Messages','source_rows':[n]})
            if text in {'DEV','LOG'}:
                kind='dev_keys' if text=='DEV' else 'log_keys';roles=[kind]
                result[kind].append({'key':row['cells']['B'],'value':row['cells']['C'],'meaning':'Client reference',
                    'usage':'In-game shortcut' if text=='DEV' else 'String identifier, not executable','scope':'Messages','source_rows':[n]})
            if text=='DEFECTS':roles=['report'];result['report_headers'].append(n)
            result['rows'].append({'row':n,'disposition':'case' if ids else 'context','case_ids':ids,'roles':roles})
        if self.missing:result['rows']=result['rows'][:-1]
        return result


def mixed_file(path,with_report=True):
    w=Workbook();s=w.active;s.title='自由形式'
    s.append(['GOAL','Messages']);s.append(['GUIDE','Use the supplied shortcut to open Messages'])
    s.cell(25,1,'TERM');s.cell(25,2,'Inbox');s.cell(25,3,'Posteingang')
    s.cell(40,1,'DEV');s.cell(40,2,'OPEN_MAIL');s.cell(40,3,'Messages shortcut')
    s.cell(55,1,'LOG');s.cell(55,2,'LOC_MAIL');s.cell(55,3,'Posteingang')
    if with_report:
        s.cell(70,1,'DEFECTS');s.cell(70,2,'Observed / why / fix');s.cell(71,2,'Example only')
    w.save(path)


def run_fixture(root,reports):
    Image.new('RGB',(480,800),'white').save(root/'screen.png')
    tax={'name_zh':'','name_en':'Typo','priority':'Minor','category':'Text','owner':''}
    project={'language':'German','cases':[], 'client_workbook':{'template':str(root/'client-template.xlsx'),'sheets':reports,
        'routes':[],'pass_sheet':reports[0]['id']}}
    state={'started_at':'2026-09-20','project':project,'cases':{},'evidence':{'E':{'path':'screen.png'}}}
    for index,report in enumerate(reports):
        cid='C'+str(index);tid='T'+str(index);project['cases'].append({'id':cid,'title':'Messages'})
        project['client_workbook']['routes'].append({'taxonomy_id':tid,'sheet':report['id']})
        findings=[{'id':cid+'-B'+str(n),'taxonomy_id':tid,'taxonomy':tax,'observed':'Posteingaang','expected':'Posteingang',
            'evidence_id':'E','bbox':{'x':.1,'y':.2,'width':.5,'height':.1},'verification':{'verdict':'confirmed','explanation':'Duplicate letter.'}} for n in range(2)]
        state['cases'][cid]={'status':'BUG','coverage':'complete','findings':findings}
    return state


class AdaptiveWorkbookTests(unittest.TestCase):
    def test_mixed_sheet_late_header_and_keys_reach_final_context(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);mixed_file(root/'client.xlsx');model=Planner([spec()])
            out=import_project(root/'client.xlsx',root/'project',{'language':'German','package':'com.demo.game'},model=model)
            p=read_json(out)
            self.assertEqual(p['import_summary']['cases'],1);self.assertTrue(p['import_summary']['ready'])
            self.assertEqual(p['glossary'][0]['target'],'Posteingang')
            self.assertEqual(p['project_context']['dev_keys'][0]['key'],'OPEN_MAIL')
            self.assertEqual(p['project_context']['log_keys'][0]['key'],'LOC_MAIL')
            self.assertIn(70,model.layout_input['sheets'][0]['report_headers'])
            self.assertIn('report',model.layout_input['sheets'][0]['roles']);self.assertIn('cases',model.layout_input['sheets'][0]['roles'])
            self.assertIn('LOC_MAIL',str(model.context_input))
            self.assertEqual(p['project_context']['case_context']['W0001']['log_keys'],['log_keys-1'])
            self.assertEqual(p['client_workbook']['sheets'][0]['columns'][0]['field'],'description')

    def test_objectives_only_uses_simple_fallback_with_optional_keys(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);mixed_file(root/'client.xlsx',False)
            p=read_json(import_project(root/'client.xlsx',root/'out',{'language':'German','package':'com.demo.game'},model=Planner()))
            self.assertTrue(p['client_workbook']['fallback'])
            w=load_workbook(root/'out/client-template.xlsx');self.assertEqual(w.sheetnames,['Bug list']);self.assertEqual(w.active.max_column,9);w.close()

    def test_missing_essential_context_cannot_start_or_leave_a_runnable_plan(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);mixed_file(root/'client.xlsx',False)
            with self.assertRaisesRegex(PilotError,'Essential guide'):
                import_project(root/'client.xlsx',root/'out',{'language':'German','package':'com.demo.game'},model=Planner(blocked=True))
            self.assertFalse((root/'out/project.json').exists())

    def test_comments_links_and_merge_hierarchy_are_visible_to_discovery(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);w=Workbook();s=w.active;s['A1']='Messages';s.merge_cells('A1:A3');s['B3']='Read all tabs'
            s['B3'].comment=Comment('Skip the tutorial','Client');s['B3'].hyperlink='https://example.invalid/guide'
            w.save(root/'c.xlsx');snap=inspect_workbook(root/'c.xlsx',root)
            row=next(r for r in snap['sheets'][0]['rows'] if r['row']==3)
            self.assertEqual(row['merged_context']['A1:A3'],'Messages')
            self.assertIn('Skip the tutorial',row['annotations']['B']['comment']);self.assertIn('/guide',row['annotations']['B']['hyperlink'])

    def test_merged_records_expand_and_preserve_footer(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);w=Workbook();s=w.active;s.title='Form'
            s.append(['Issue','Details']);s['A2']='Text / issue / fix';s['A3']='Screenshot';s.merge_cells('B2:C2');s.merge_cells('B3:C3')
            s['A4']='Client footer';w.save(root/'client-template.xlsx')
            r=spec('Form',2,2);r['columns'].append({'column':2,'row_offset':1,'field':'screenshot','literal':''})
            out=export_client(root,state=run_fixture(root,[r]));b=load_workbook(out);sh=b.active
            self.assertEqual(sh['A6'].value,'Client footer');self.assertEqual(sh['A4'].value,'Text / issue / fix')
            self.assertTrue({'B2:C2','B3:C3','B4:C4','B5:C5'}<=set(map(str,sh.merged_cells.ranges)))
            self.assertEqual(len(sh._images),2);self.assertIn('Posteingang',sh['B4'].value);b.close()

    def test_multiple_report_regions_on_one_sheet_keep_both_and_unique_picture_ids(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);w=Workbook();s=w.active;s.title='Reports'
            s['A1']='Text';s['B2']='Example';s['A4']='Visual';s['B5']='Example';s['A6']='Footer';w.save(root/'client-template.xlsx')
            one=spec('Reports',2);two=spec('Reports',5);two['id']='visual'
            b=load_workbook(export_client(root,state=run_fixture(root,[one,two])))
            self.assertEqual(b.sheetnames,['Reports']);self.assertEqual(b.active['A5'].value,'Visual');self.assertEqual(b.active['A8'].value,'Footer')
            self.assertEqual(len(b.active._images),4);b.close()

    def test_report_example_is_not_a_case_and_pass_has_no_image_or_comment(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);w=Workbook();s=w.active;s.title='Notes';s.append(['Defect']);s.append(['Example']);w.save(root/'client-template.xlsx')
            r=spec('Notes',2);r['columns'][0]['column']=1
            state=run_fixture(root,[r]);state['cases']['C0']={'status':'PASS','coverage':'complete','findings':[]}
            b=load_workbook(export_client(root,state=state));self.assertEqual(b.active['A2'].value,'Pass');self.assertFalse(b.active._images);b.close()

    def test_import_failure_never_constructs_android_or_engine(self):
        with tempfile.TemporaryDirectory() as t,patch('lqa_pilot.workbook_project.import_project',side_effect=PilotError('Ambiguous scope')),patch('lqa_pilot.engine.Engine') as engine,patch('lqa_pilot.android.Android') as android:
            jobs=Jobs(t);jobs.current={'id':'fixture','status':'importing','downloads':[]};jobs.work(Path(t)/'fixture',{'mode':'full'})
            self.assertEqual(jobs.current['status'],'error');engine.assert_not_called();android.assert_not_called()
