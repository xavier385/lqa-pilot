"""Run against an already exported offline demo; do not require GPT or Android."""
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from lqa_pilot.project import load_project
from lqa_pilot.recheck import prepare_recheck
from lqa_pilot.report import CLIENT_HEADERS
from lqa_pilot.workbook_meta import read_metadata

folder = Path(sys.argv[1]).resolve()
wb = load_workbook(folder/"report.xlsx")
assert wb.sheetnames == ["Bug list"]
sheet = wb.active
assert [c.value for c in sheet[1]] == CLIENT_HEADERS
assert len(sheet._images) == 1
assert sheet["P2"].value == "Pass"
assert all(sheet.cell(2, i).value in {None, ""} for i in range(1,19) if i not in {2,3,16})
assert len(sheet.data_validations.dataValidation) == 1
assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
picture = sheet._images[0]
assert picture.anchor._from.row == 2 and picture.anchor._from.col == 3
assert picture.width == 800 and picture.height == 1100

# Simulate an operator selecting one row in Excel, preserving all other ZIP parts.
ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
selection = folder/"review-selection-test.xlsx"
with ZipFile(folder/"report.xlsx") as source, ZipFile(selection, "w", ZIP_DEFLATED) as target:
    part = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
    row = next(r for r in part.iter(f"{{{ns}}}row") if r.attrib["r"] == "2")
    cell = next(c for c in row if c.attrib["r"] == "P2")
    cell.clear()
    cell.attrib = {"r":"P2", "t":"inlineStr"}
    ET.SubElement(ET.SubElement(cell, f"{{{ns}}}is"), f"{{{ns}}}t").text = "Recheck"
    for name in source.namelist():
        target.writestr(name, ET.tostring(part) if name == "xl/worksheets/sheet1.xml" else source.read(name))
output = Path(sys.argv[2]).resolve()
project = load_project(prepare_recheck(selection, folder, output))
assert [c["id"] for c in project["cases"]] == ["DEMO-PASS"]
assert len(project["review_context"]) == 1
# A normal spreadsheet save must retain the invisible identity used for rechecks.
saved = folder/"roundtrip.xlsx"
wb.save(saved)
assert read_metadata(saved) == read_metadata(folder/"report.xlsx")
print("Excel verified: one sheet, 18 client columns, original-resolution image in Screenshot, blank Pass fields, Recheck import, metadata roundtrip.")
