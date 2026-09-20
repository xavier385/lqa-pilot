"""Local fixtures only: no GPT, Android, account login or cloud deployment."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
from zipfile import ZipFile

from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from PIL import Image

from lqa_pilot.client_excel import export_client,xml
from lqa_pilot.workbook_project import import_project,inspect_workbook
from lqa_pilot.web_bridge import Companion,Jobs
from lqa_pilot.project import load_project
from lqa_pilot.util import PilotError,read_json

HEADERS=['Where','Case','Picture','Kind','Text / issue / fix','Key','Priority','Status']
FIELDS=['location','case_id','screenshot','type','description','id','priority','status']

def fixture(path,two=True,case_count=2):
    # A deliberately small synthetic client input, not a report for the user.
    w=Workbook();s=w.active;s.title='Journeys';s.append(['Goal'])
    for i in range(case_count):s.append([f'Open menu {i+1} and read all tabs'])
    for name in ['Text bugs','Visual bugs'] if two else ['Text bugs']:
        s=w.create_sheet(name);s.append(HEADERS);s.append(['Example']*8)
        s.row_dimensions[1].height=40;s.row_dimensions[2].height=200
        s.freeze_panes='C2';s.column_dimensions['E'].width=55
        for c in s[1]:c.font=Font(name='Calibri',size=12,bold=True,color='123456');c.fill=PatternFill('solid',fgColor='AABBCC')
        for c in s[2]:c.font=Font(name='Calibri',size=10)
    w.save(path)

def mapping(name):
    return {'name':name,'role':'bugs','header_row':1,'data_start':2,'data_end':2,'style_row':2,
            'purpose':name,'columns':[{'column':i+1,'field':f} for i,f in enumerate(FIELDS)]}

class FakePlanner:
    def __init__(self,two=True,omit=False):self.two,self.omit,self.read_rows=two,omit,[]
    def ask(self,role,prompt,schema,images=()):
        self.before_call(role)
        if prompt.startswith('Map output'):
            reports=[]
            for name in ['Text bugs','Visual bugs'] if self.two else ['Text bugs']:
                r=mapping(name);r.pop('role');r.update(id=name,row_span=1)
                r['columns']=[dict(c,row_offset=0,literal='') for c in r['columns']];reports.append(r)
            return {'reports':reports,'notes':[],'blockers':[]}
        if prompt.startswith('Review the plan'):
            data=json.loads(prompt.split('DATA:\n')[1])
            return {'ready':True,'issues':[],'context_summary':'Synthetic test fixture',
                    'cases':[{'id':c['id'],'guide':'Read all tabs','dev_keys':[],'log_keys':[]} for c in data['cases']]}
        if prompt.startswith('Route'):
            data=json.loads(prompt.split('DATA:\n')[1])
            return {'routes':[{'taxonomy_id':t['id'],'sheet':'Text bugs'} for t in data['taxonomy']['entries']], 'pass_sheet':'Text bugs'}
        data=json.loads(prompt.split('DATA:\n')[1]);rows=[];cases=[]
        if data['sheet']!='Journeys':
            return {'rows':[{'row':r['row'],'disposition':'context','case_ids':[],'roles':['report']} for r in data['rows']],
                'cases':[],'instructions':[],'glossary':[],'cheats':[],'dev_keys':[],'log_keys':[],
                'report_headers':[1] if any(r['row']==1 for r in data['rows']) else [],'taxonomy':[],'notes':[],'continuation':'Bug table'}
        for r in data['rows']:
            n=r['row'];self.read_rows.append(n)
            ids=[]
            if n>1:
                ident='ROW-'+str(n);ids=[ident]
                cases.append({'id':ident,'title':'Menu '+str(n),'objective':r['cells']['A'],'journey':['Discover the requested menu'],
                    'checks':[{'id':'C'+str(n),'scope':'All tabs','instruction':'Read all visible text','modality':'visual'}],'source_rows':[n]})
            rows.append({'row':n,'disposition':'case' if ids else 'context','case_ids':ids,'roles':['cases']})
        if self.omit:rows=rows[:-1]
        return {'rows':rows,'cases':cases,'instructions':[],'glossary':[],'cheats':[],'taxonomy':[],'continuation':'',
                'dev_keys':[],'log_keys':[],'report_headers':[],'notes':[]}


class WorkbookTests(unittest.TestCase):
    def test_excel_compatibility_prefixes_survive_serialization(self):
        import xml.etree.ElementTree as ET
        source=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac"><sheetData/></worksheet>'
        result=xml(ET.fromstring(source),source)
        ET.fromstring(result)
        self.assertIn(b'xmlns:x14ac=',result)
        self.assertIn(b'mc:Ignorable="x14ac"',result)

    def test_import_reads_rows_beyond_first_batch_and_keeps_provenance(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);fixture(root/'client.xlsx',case_count=75);model=FakePlanner()
            file=import_project(root/'client.xlsx',root/'project',{'language':'German','package':'com.demo.game','mode':'fast'},model=model)
            p=load_project(file)
            self.assertEqual(len(p['cases']),75);self.assertEqual(model.read_rows,list(range(1,77)))
            self.assertEqual(p['planning_notes']['source_rows']['W0075']['rows'],[76])
            self.assertEqual(len(p['client_workbook']['sheets']),2)
            self.assertEqual(p['mode'],'fast')

    def test_missing_row_cannot_start_a_partial_plan(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);fixture(root/'client.xlsx')
            with self.assertRaises(PilotError):
                import_project(root/'client.xlsx',root/'project',{'language':'German','package':'com.demo.game'},model=FakePlanner(omit=True))
            self.assertFalse((root/'project/project.json').exists())

    def test_client_output_keeps_two_sheets_styles_and_routes_images(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);fixture(root/'client-template.xlsx')
            Image.new('RGB',(800,1100),'white').save(root/'screen.png')
            tax={'name_zh':'','name_en':'Spelling','priority':'T2','category':'Text','owner':''}
            finding={'id':'B1','taxonomy_id':'text','taxonomy':tax,'check_id':'C','evidence_id':'E','bbox':{'x':.1,'y':.2,'width':.5,'height':.1},
                     'observed':'=Untrusted game text','expected':'Correct text','verification':{'verdict':'confirmed','explanation':'Reason','report_comment':'A letter is missing.'}}
            visual=copy.deepcopy(finding);visual.update(id='B2',taxonomy_id='visual')
            contract={'template':str(root/'client-template.xlsx'),'sheets':[mapping('Text bugs'),mapping('Visual bugs')],
                      'routes':[{'taxonomy_id':'text','sheet':'Text bugs'},{'taxonomy_id':'visual','sheet':'Visual bugs'}],'pass_sheet':'Text bugs'}
            run={'started_at':'2026-09-19','project':{'client_workbook':contract,'cases':[{'id':'A','title':'Menu A'},{'id':'B','title':'Menu B'}]},
                 'cases':{'A':{'status':'BUG','coverage':'complete','findings':[finding,visual]},'B':{'status':'PASS','coverage':'complete','findings':[]}},
                 'evidence':{'E':{'path':'screen.png'}}}
            output=export_client(root,state=run)
            with ZipFile(root/'client-template.xlsx') as src,ZipFile(output) as dst:
                self.assertEqual(src.read('xl/styles.xml'),dst.read('xl/styles.xml'))
            w=load_workbook(output)
            self.assertEqual(w.sheetnames,['Text bugs','Visual bugs'])
            for s in w:
                self.assertEqual([c.value for c in s[1]],HEADERS)
                self.assertEqual(s['A1'].fill.fgColor.rgb,'00AABBCC')
                self.assertEqual(s.column_dimensions['E'].width,55)
                self.assertEqual(len(s._images),1)
                self.assertEqual((s._images[0].width,s._images[0].height),(800,1100))
                self.assertEqual(s['E2'].data_type,'s')
            self.assertEqual(w['Text bugs']['H3'].value,'Pass')
            self.assertTrue(all(w['Text bugs'].cell(3,c).value in {None,''} for c in (3,4,5,6,7)))
            self.assertEqual(sorted(p.name for p in root.glob('*.xlsx')),['client-template.xlsx','report.xlsx'])


class BridgeTests(unittest.TestCase):
    def test_pair_is_origin_bound_and_rejects_wrong_token(self):
        with tempfile.TemporaryDirectory() as t:
            c=Companion(t,'https://qa.example.com');code=c.code
            with self.assertRaises(PilotError):c.pair('https://evil.example',code)
            token=c.pair('https://qa.example.com',code)
            self.assertTrue(c.authorized('https://qa.example.com','Bearer '+token))
            self.assertFalse(c.authorized('https://evil.example','Bearer '+token))
            self.assertFalse(c.authorized('https://qa.example.com','Bearer wrong'))

    def test_submit_validates_mode_before_starting_worker(self):
        with tempfile.TemporaryDirectory() as t:
            j=Jobs(t)
            with self.assertRaises(PilotError):j.submit(b'file','input.xlsx',{'mode':'unknown'})
            self.assertIsNone(j.current)

    def test_stop_during_import_sets_cancellation_without_android(self):
        with tempfile.TemporaryDirectory() as t:
            j=Jobs(t);j.current={'id':'safe-fixture','status':'importing','downloads':[]}
            j.stop();self.assertTrue(j.cancelled.is_set());self.assertEqual(j.current['status'],'stopping')

    def test_download_cannot_select_an_arbitrary_file(self):
        with tempfile.TemporaryDirectory() as t:
            j=Jobs(t);j.current={'id':'safe-fixture','status':'complete','downloads':[]}
            with self.assertRaises(PilotError):j.download('../../auth')
