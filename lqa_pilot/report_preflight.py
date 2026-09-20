"""Private synthetic export check. Never exposed as a game result or downloadable report."""
import copy
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
from .client_excel import export_client,report_id


def check_export(folder,project,taxonomy):
    folder=Path(folder);probe=folder/'export-preflight';probe.mkdir(exist_ok=True)
    Image.new('RGB',(480,800),'white').save(probe/'synthetic.png')
    sample=copy.deepcopy(project);sample['cases']=[]
    contract=sample['client_workbook'];contract['template']=str(folder/'client-template.xlsx');contract['routes']=[]
    state={'started_at':'2000-01-01','project':sample,'cases':{},'evidence':{'fixture':{'path':'synthetic.png'}}}
    for index,spec in enumerate(contract['sheets']):
        ident='PREFLIGHT-'+str(index);key='TYPE-'+str(index)
        contract['routes'].append({'taxonomy_id':key,'sheet':report_id(spec)})
        sample['cases'].append({'id':ident,'title':'SOFTWARE PREFLIGHT, NOT GAME EVIDENCE'})
        finding={'id':ident,'taxonomy_id':key,'taxonomy':taxonomy['entries'][0],
            'observed':'Synthetic export check','expected':'Synthetic correction','evidence_id':'fixture',
            'bbox':{'x':.1,'y':.2,'width':.5,'height':.1},
            'verification':{'verdict':'confirmed','explanation':'Synthetic layout check'}}
        second=copy.deepcopy(finding);second['id']+='-2'
        state['cases'][ident]={'status':'BUG','coverage':'complete','findings':[finding,second]}
    output=export_client(probe,state=state)
    book=load_workbook(output)
    if set(book.sheetnames)!={s['name'] for s in contract['sheets']}:raise RuntimeError('Preflight: fogli report non corrispondenti')
    book.close()
