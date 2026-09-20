"""Read client workbooks into finite plans; preserve row-level provenance."""
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from .util import PilotError, digest

def inspect_workbook(path, folder):
    path, folder = Path(path), Path(folder)
    if path.stat().st_size > 30*1024*1024:
        raise PilotError("Excel troppo grande: limite 30 MB")
    try:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 180*1024*1024 or len(z.infolist()) > 8000:
                raise PilotError("Excel oltre i limiti di elaborazione")
            if any('vbaProject' in n or '/embeddings/' in n for n in z.namelist()):
                raise PilotError("Usare un .xlsx senza macro o oggetti eseguibili incorporati")
        wb = load_workbook(path, data_only=True)
        formulas = load_workbook(path, data_only=False, read_only=True)
    except (ValueError, zipfile.BadZipFile, OSError) as e:
        raise PilotError("File Excel non valido o protetto da password") from e
    sheets, images, cells = [], [], 0
    if len(wb.worksheets) > 40:
        raise PilotError("Suddividere il progetto: massimo 40 fogli per Excel")
    for index, sh in enumerate(wb):
        if sh.max_row > 25000 or sh.max_column > 200:
            raise PilotError(f"Foglio {sh.title}: area troppo estesa; rimuovere righe/colonne inutilizzate")
        rows = []
        for row in sh:
            values = {c.column_letter:str(c.value) for c in row if c.value is not None}
            annotations={c.column_letter:{'comment':c.comment.text if c.comment else '',
                'hyperlink':(c.hyperlink.target or c.hyperlink.location) if c.hyperlink else ''}
                for c in row if getattr(c,'comment',None) or getattr(c,'hyperlink',None)}
            if values or annotations:
                rows.append({"row":row[0].row, "cells":values,'annotations':annotations})
                cells += len(values)
        for row in formulas[sh.title]:
            for c in row:
                if c.data_type == 'f' and sh[c.coordinate].value is None:
                    current=next((r for r in rows if r['row']==c.row),None)
                    if current is None:
                        current={'row':c.row,'cells':{}};rows.append(current)
                    current['cells'][c.column_letter]='[Formula without cached result] '+str(c.value)
        if cells > 45000:
            raise PilotError("Troppi contenuti: suddividere il progetto, senza omettere test case")
        for n, picture in enumerate(sh._images):
            target = folder/'references'/f'sheet-{index}-image-{n}.{picture.format}'
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(picture._data())
            anchor = getattr(picture.anchor, '_from', None)
            images.append({"sheet":sh.title,"row":anchor.row+1 if anchor else 1,"path":str(target.resolve())})
            image_row=anchor.row+1 if anchor else 1
            if not any(r['row']==image_row for r in rows):
                rows.append({'row':image_row,'cells':{'[image]':'Embedded reference image at this row'}})
        rows.sort(key=lambda r:r['row'])
        for row in rows:
            row['merged_context']={str(region):str(sh.cell(region.min_row,region.min_col).value)
                for region in sh.merged_cells.ranges if region.min_row<=row['row']<=region.max_row
                and sh.cell(region.min_row,region.min_col).value is not None}
        sheets.append({"name":sh.title,"max_row":sh.max_row,"max_column":sh.max_column,"rows":rows,
                       "merged_cells":[str(r) for r in sh.merged_cells.ranges],
                       'visibility':sh.sheet_state,'tables':[{ 'name':t.name,'range':t.ref} for t in sh.tables.values()]})
    wb.close(); formulas.close()
    return {"sheets":sheets,"images":images,"sha256":digest(path.read_bytes())}


from .workbook_understanding import import_project
